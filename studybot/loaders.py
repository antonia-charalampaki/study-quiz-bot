"""
Φόρτωση σημειώσεων από αρχεία .txt, .md, .pdf και .docx.
"""

from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document

SUPPORTED = (".txt", ".md", ".pdf", ".docx")
MIN_CHARS = 200          # πολύ λίγο κείμενο δεν φτάνει για ερωτήσεις
MAX_CHARS = 200_000      # περίπου 80-100 σελίδες, για να μην ξεφύγει το κόστος/χρόνος


class NotesError(Exception):
    """Φιλικό σφάλμα για προβλήματα με το αρχείο σημειώσεων."""


def _read_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _read_docx(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # και τα κείμενα μέσα σε πίνακες
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Παίρνει το όνομα και τα bytes ενός αρχείου και επιστρέφει καθαρό κείμενο.
    (Με bytes, ώστε να δουλεύει και με αρχεία που ανεβαίνουν από το web.)"""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise NotesError(f"Το αρχείο {filename} δεν υποστηρίζεται. Δεκτά: {', '.join(SUPPORTED)}")

    if suffix == ".pdf":
        text = _read_pdf(data)
    elif suffix == ".docx":
        text = _read_docx(data)
    else:
        text = data.decode("utf-8", errors="replace")

    text = text.strip()
    if len(text) < MIN_CHARS:
        if suffix == ".pdf":
            raise NotesError("Δεν βρέθηκε αρκετό κείμενο στο PDF. Μήπως είναι σκαναρισμένο (εικόνες);")
        raise NotesError("Οι σημειώσεις είναι πολύ λίγες. Βάλε τουλάχιστον μία παράγραφο.")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
    return text


def load_notes(path: Path) -> str:
    """Διαβάζει σημειώσεις από αρχείο στον υπολογιστή."""
    if not path.exists():
        raise NotesError(f"Δεν βρέθηκε το αρχείο: {path}")
    return extract_text(path.name, path.read_bytes())
